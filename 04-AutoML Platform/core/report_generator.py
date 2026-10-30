"""报告生成器：输出 Markdown 与 Word（纯文字/表格，不含图片）。"""

import json
from pathlib import Path
from typing import Dict, List, Optional

import pandas as pd


class ReportGenerator:
    def __init__(self, experiment_name: str, output_dir: str):
        self.experiment_name = experiment_name
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.sections: List[str] = []

    def add_header(self, title: str, level: int = 1):
        self.sections.append(f"{'#' * level} {title}\n")

    def add_text(self, text: str):
        self.sections.append(text + "\n")

    def add_table(self, df: pd.DataFrame, title: str = None):
        if title:
            self.sections.append(f"**{title}**\n")
        self.sections.append(df.to_markdown(index=False) + "\n\n")

    def to_markdown(self) -> str:
        return "\n".join(self.sections)

    def save_markdown(self, filename: str = "report.md"):
        path = self.output_dir / filename
        path.write_text(self.to_markdown(), encoding="utf-8")
        return str(path)

    def save_word(self, filename: str = "report.docx"):
        """将 Markdown 内容导出为 Word 文档（python-docx）。"""
        try:
            from docx import Document
            from docx.shared import Inches, Pt
            from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
        except ImportError:
            raise ImportError("Word 导出需要 python-docx，请执行：pip install python-docx")

        doc = Document()
        # 标题
        title = doc.add_heading(f"实验报告：{self.experiment_name}", level=0)
        title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

        # 按章节写入
        for section in self.sections:
            section = section.rstrip("\n")
            if not section:
                continue
            if section.startswith("#"):
                level = len(section) - len(section.lstrip("#"))
                text = section.lstrip("#").strip()
                doc.add_heading(text, level=min(level, 4))
            elif section.startswith("**") and section.endswith("**"):
                # 表格标题
                p = doc.add_paragraph()
                run = p.add_run(section.strip("*"))
                run.bold = True
                run.font.size = Pt(11)
            elif "|" in section and "\n" in section:
                # Markdown 表格
                self._add_markdown_table_to_doc(doc, section)
            else:
                doc.add_paragraph(section)

        path = self.output_dir / filename
        doc.save(str(path))
        return str(path)

    @staticmethod
    def _add_markdown_table_to_doc(doc, md_table: str):
        """把 Markdown 表格字符串写入 Word。"""
        lines = [ln.strip() for ln in md_table.split("\n") if ln.strip()]
        if len(lines) < 2:
            return
        # 第一行表头，第二行分隔线，后续数据
        headers = [c.strip() for c in lines[0].split("|") if c.strip()]
        data_lines = lines[2:] if len(lines) >= 3 else lines[1:]
        rows = []
        for ln in data_lines:
            cells = [c.strip() for c in ln.split("|") if c.strip()]
            if cells:
                rows.append(cells)

        if not headers:
            return
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = "Light Grid Accent 1"
        hdr_cells = table.rows[0].cells
        for i, h in enumerate(headers):
            hdr_cells[i].text = h
        for row in rows:
            row_cells = table.add_row().cells
            for i, cell in enumerate(row[:len(headers)]):
                row_cells[i].text = cell

    #: 可选章节（顺序即报告中的排列顺序）
    DEFAULT_SECTIONS = [
        "overview", "preprocessing", "feature_selection", "base_models",
        "ensemble", "transformer", "automl", "survival",
        "shap", "ablation", "minimal_subset", "conclusion",
    ]

    _CN_NUMS = ["一", "二", "三", "四", "五", "六", "七", "八", "九", "十",
                "十一", "十二", "十三", "十四", "十五", "十六"]

    def _t(self, name: str) -> str:
        """生成动态编号的中文章节标题。"""
        self._sec_no += 1
        num = self._CN_NUMS[self._sec_no - 1] if self._sec_no <= len(self._CN_NUMS) else str(self._sec_no)
        return f"{num}、{name}"

    def build_from_results(
        self,
        config: Dict,
        leaderboard: pd.DataFrame,
        shap_indicator: Optional[str] = None,
        sections: Optional[Dict[str, bool]] = None,
    ):
        if sections is None:
            sections = {k: True for k in self.DEFAULT_SECTIONS}
        self._sec_no = 0

        self.add_header(f"实验报告：{self.experiment_name}")
        self.add_text("本报告由 AutoML-Platform 自动生成，汇总实验结果。报告仅包含文字与表格，便于直接粘贴到论文或补充材料中。")

        if sections.get("overview", True):
            self.add_header(self._t("实验概述"), level=2)
            for k, v in config.items():
                if k in {"preprocessing", "feature_selection", "modeling", "ensemble",
                         "transformer", "automl", "survival", "ablation", "minimal_subset"}:
                    continue
                self.add_text(f"- **{k}**: {v}")

        if sections.get("preprocessing", True):
            self._add_preprocessing_section(config.get("preprocessing"))

        if sections.get("feature_selection", True):
            self._add_feature_selection_section(config.get("feature_selection"))

        if sections.get("base_models", True):
            self._add_base_models_section(config, leaderboard)

        if sections.get("ensemble", True):
            self._add_module_section(self._t("集成学习结果"), config.get("ensemble"), "ensemble")

        if sections.get("transformer", True):
            self._add_module_section(self._t("Transformer 结果"), config.get("transformer"), "transformer")

        if sections.get("automl", True):
            self._add_module_section(self._t("AutoML 结果"), config.get("automl"), "automl")

        if sections.get("survival", True):
            self._add_module_section(self._t("生存分析结果"), config.get("survival"), "survival")

        if sections.get("shap", True):
            self._add_shap_section(config, shap_indicator)

        if sections.get("ablation", True):
            self._add_ablation_section(config)

        if sections.get("minimal_subset", True):
            self._add_minimal_subset_section(config)

        if sections.get("conclusion", True):
            self._add_conclusion(config, shap_indicator, sections)

        return self.to_markdown()

    def _add_base_models_section(self, config: Dict, leaderboard: pd.DataFrame):
        self.add_header(self._t("基础模型训练结果"), level=2)
        modeling_config = config.get("modeling") or {}
        leaderboard_path = modeling_config.get("leaderboard_path")
        if leaderboard_path and Path(leaderboard_path).exists():
            df_lb = pd.read_csv(leaderboard_path)
            self.add_table(df_lb, title="基础模型完整评估结果（OOF）")
            self.add_text("表格包含 AUC、AUPRC、准确率、敏感度、特异度、PPV、NPV、F1、Kappa、LogLoss、Brier Score 等 11 项指标。")
        elif leaderboard is not None and not leaderboard.empty:
            self.add_table(leaderboard, title="基础模型评估结果（OOF）")
        else:
            self.add_text("暂无基础模型结果。")

    def _add_ablation_section(self, config: Dict):
        """消融分析章节：变量榜单 + ΔAUC 矩阵 + 推荐回加变量。"""
        self.add_header(self._t("消融分析结果"), level=2)
        ab_cfg = config.get("ablation") or {}
        result_path = ab_cfg.get("result_path")
        if not result_path or not Path(result_path).exists():
            self.add_text("暂无消融分析结果。")
            return
        try:
            from core.ablation import summarize_variables, make_delta_matrix
            df = pd.read_csv(result_path)
            n_ok = int((df["status"] == "ok").sum())
            n_fail = int((df["status"] == "failed").sum())
            self.add_text(f"- 消融组合：共 {len(df)} 个（成功 {n_ok}，失败 {n_fail}）。")
            summary = summarize_variables(df)
            if not summary.empty:
                self.add_table(summary, title="消融变量榜单（按平均 ΔAUC 排序）")
                rec = summary[summary["recommended"] == True]
                if len(rec):
                    self.add_text(f"- 达到推荐回加标准的变量：{', '.join(rec['variable'].astype(str).tolist())}")
                else:
                    self.add_text("- 无变量达到推荐回加标准。")
            mat = make_delta_matrix(df)
            if not mat.empty:
                self.add_table(mat.round(4).reset_index().rename(columns={"index": "variable"}),
                               title="消融 ΔAUC 矩阵（变量 × 模型）")
        except Exception as e:
            self.add_text(f"消融结果读取失败：{e}")

    def _add_minimal_subset_section(self, config: Dict):
        """最小子集分析章节：各模型最小子集对比 + 子集特征清单。"""
        self.add_header(self._t("最小协变量子集分析"), level=2)
        ms_cfg = config.get("minimal_subset") or {}
        summary_path = ms_cfg.get("summary_path")
        if not summary_path or not Path(summary_path).exists():
            self.add_text("暂无最小子集分析结果。")
            return
        try:
            df = pd.read_csv(summary_path)
            self.add_table(df, title="最小子集对比汇总（子集大小 / AUC / ΔAUC / DeLong p）")
            ms_dir = Path(summary_path).parent
            lines = []
            for _, r in df[df["status"] == "ok"].iterrows():
                subset_path = ms_dir / f"{r['model']}_subset_features.csv"
                if subset_path.exists():
                    feats = pd.read_csv(subset_path)["feature"].astype(str).tolist()
                    better = "（严格更优）" if r.get("strictly_better") else ""
                    lines.append(f"- **{r['model']}**（{r['n_subset']} 个特征{better}）：{', '.join(feats)}")
            if lines:
                self.add_text("各模型最小子集特征清单：")
                for ln in lines:
                    self.add_text(ln)
        except Exception as e:
            self.add_text(f"最小子集结果读取失败：{e}")

    def _add_conclusion(self, config: Dict, shap_indicator: Optional[str], sections: Dict[str, bool]):
        self.add_header(self._t("结论与建议"), level=2)
        included = [k for k, v in sections.items() if v]
        self.add_text(f"- 本报告包含以下部分：{', '.join(included)}。")
        self.add_text("- 建议结合可视化选项卡中的图片（ROC、SHAP 图等）对模型性能与特征重要性进行进一步解读。")
        if shap_indicator:
            self.add_text(f"- 重点关注指标 **{shap_indicator}** 的跨模型 SHAP 分析结果，可作为领域/业务解释的主要依据。")

    def _add_preprocessing_section(self, prep_config):
        """添加预处理说明章节。"""
        self.add_header(self._t("预处理"), level=2)
        if not prep_config:
            self.add_text("暂无预处理结果。")
            return

        self.add_text(f"- **缺失值策略**：{prep_config.get('missing_strategy')}")
        self.add_text(f"- **异常值处理**：{prep_config.get('outlier_method')}")
        self.add_text(f"- **编码方式**：{prep_config.get('encoding')}")
        self.add_text(f"- **缩放方式**：{prep_config.get('scaling')}")
        self.add_text(f"- **不平衡处理**：{prep_config.get('imbalance_method')}")

        report_path = prep_config.get("report_path")
        if report_path and Path(report_path).exists():
            with open(report_path, "r", encoding="utf-8") as f:
                report = json.load(f)
            self.add_text(f"- **连续变量数**：{report.get('n_numeric')}")
            self.add_text(f"- **分类变量数**：{report.get('n_categorical')}")
            self.add_text(f"- **透传变量数**：{report.get('n_passthrough')}")
            self.add_text(f"- **输出数据形状**：{report.get('output_shape')}")
            self.add_text(f"- **类别分布**：{report.get('class_distribution')}")
            if report.get("numeric_columns"):
                self.add_text(f"- **连续变量**：{', '.join(report.get('numeric_columns', []))}")
            if report.get("categorical_columns"):
                self.add_text(f"- **分类变量**：{', '.join(report.get('categorical_columns', []))}")
            if report.get("passthrough_columns"):
                self.add_text(f"- **透传变量**：{', '.join(report.get('passthrough_columns', []))}")

    def _add_feature_selection_section(self, fs_config):
        """添加特征选择结果章节。"""
        self.add_header(self._t("特征选择结果"), level=2)
        if not fs_config:
            self.add_text("暂无特征选择结果。")
            return

        if isinstance(fs_config, dict):
            self.add_text(f"- **方法**：{', '.join(fs_config.get('methods', []))}")
            self.add_text(f"- **共识阈值**：{fs_config.get('consensus_threshold')}")
            self.add_text(f"- **随机种子**：{fs_config.get('random_state')}")

            consensus_path = fs_config.get("consensus_table_path")
            if consensus_path and Path(consensus_path).exists():
                df_cons = pd.read_csv(consensus_path)
                n_consensus = int(df_cons["consensus"].sum()) if "consensus" in df_cons.columns else 0
                self.add_text(f"- **共识特征数量**：{n_consensus}")
                method_cols = [c for c in df_cons.columns if c not in {"feature", "n_methods_selected", "selected_methods", "consensus"}]
                if method_cols:
                    self.add_table(df_cons, title="特征选择共识表（1=选中，0=未选中）")

            summary_path = fs_config.get("feature_selection_summary_path")
            if summary_path and Path(summary_path).exists():
                df_summary = pd.read_csv(summary_path)
                self.add_table(df_summary, title="各方法选择摘要")

            stability_path = fs_config.get("stability_report_path")
            if stability_path and Path(stability_path).exists():
                df_stab = pd.read_csv(stability_path)
                self.add_table(df_stab, title="特征选择稳定性（5-fold Jaccard）")
                self.add_text("avg_jaccard 越接近 1，表示该方法选中的特征在不同数据折上越稳定。")
        elif isinstance(fs_config, list):
            self.add_text(f"- **方法**：{', '.join(fs_config)}")

    def _add_module_section(self, title: str, module_config, key: str):
        """读取模块结果 CSV/JSON 并添加到报告。"""
        self.add_header(title, level=2)
        if not module_config:
            self.add_text(f"暂无{title.replace('结果', '')}结果。")
            return

        if isinstance(module_config, dict):
            result_path = module_config.get("result_path")
            if result_path and Path(result_path).exists():
                df = pd.read_csv(result_path)
                self.add_table(df, title=f"{title}")
            else:
                self.add_text(f"暂无{title.replace('结果', '')}结果文件。")

            details_path = module_config.get("details_path")
            if details_path and Path(details_path).exists():
                with open(details_path, "r", encoding="utf-8") as f:
                    details = json.load(f)
                self.add_text(f"- **数据源**：{details.get('data_source')}")
                self.add_text(f"- **样本数**：{details.get('n_samples')}")
                self.add_text(f"- **特征数**：{details.get('n_features')}")
                if "use_optuna" in details:
                    self.add_text(f"- **Optuna 调参**：{details.get('use_optuna')}（trials={details.get('n_trials')}）")
                if "time_limit" in details:
                    self.add_text(f"- **时间限制**：{details.get('time_limit')} 秒")
                if "use_gpu" in details:
                    self.add_text(f"- **GPU**：{'启用' if details.get('use_gpu') else '未启用'}（系统 GPU 可用：{details.get('gpu_available')}）")
                if details.get("best_configs"):
                    self.add_text(f"- **最优配置**：{details.get('best_configs')}")

            params = module_config.get("params") or module_config.get("method_params")
            if params:
                self.add_text(f"- **超参数/配置**：`{params}`")
            # ensemble 等模块存的是复数 best_configs，兼容单数 best_config
            best_config = module_config.get("best_config") or module_config.get("best_configs")
            if best_config:
                self.add_text(f"- **最优配置**：`{best_config}`")
        else:
            self.add_text(f"- **配置**：{module_config}")

    def _add_shap_section(self, config: Dict, shap_indicator: Optional[str]):
        """添加 SHAP 分析章节，包含聚合排名与跨模型分析表。"""
        self.add_header(self._t("SHAP 可解释性分析"), level=2)
        modeling_config = config.get("modeling") or {}
        if not modeling_config or not modeling_config.get("output_dir"):
            self.add_text("暂无 SHAP 分析结果。")
            return

        exp_dir = Path(modeling_config["output_dir"]).parent
        shap_dir = exp_dir / "shap"
        if not shap_dir.exists():
            self.add_text("未找到 SHAP 输出目录，请先运行全模型 SHAP 分析。")
            return

        # 元信息
        meta_path = shap_dir / "shap_meta.json"
        if meta_path.exists():
            with open(meta_path, "r", encoding="utf-8") as f:
                meta = json.load(f)
            self.add_text(
                f"- 全模型 SHAP 分析：共 {meta.get('n_models_total')} 个模型，"
                f"成功 {meta.get('n_models_success')} 个，失败 {meta.get('n_models_failed')} 个。"
            )
            failures = meta.get("failures") or []
            if failures:
                self.add_text(f"- 失败模型：{', '.join(failures)}")

        # 聚合排名表
        ranking_path = shap_dir / "shap_ranking_all_models.csv"
        if ranking_path.exists():
            df_rank = pd.read_csv(ranking_path)
            self.add_table(df_rank, title="特征 SHAP 聚合排名（所有成功模型）")
            self.add_text(
                "列说明：total_mean_abs_shap 为该特征在所有模型上的 |SHAP| 之和；"
                "overall_rank 为按 total_mean_abs_shap 排序后的名次；"
                "mean_rank_across_models 为该特征在单个模型内部排名的平均值。"
            )
        else:
            self.add_text("未找到特征 SHAP 聚合排名表。")

        # 指定指标跨模型分析
        indicator = shap_indicator
        if not indicator:
            # 尝试从 meta 中读取
            if meta_path.exists():
                with open(meta_path, "r", encoding="utf-8") as f:
                    indicator = json.load(f).get("indicator")

        if indicator:
            analysis_path = shap_dir / f"shap_analysis_{indicator}.csv"
            if analysis_path.exists():
                df_analysis = pd.read_csv(analysis_path)
                self.add_table(df_analysis, title=f"选中指标 {indicator} 的跨模型 SHAP 分析")
                self.add_text(
                    "列说明：mean_abs_shap 为该指标在对应模型中的平均 |SHAP|；"
                    "rank_in_model 为该指标在模型内部的重要性排名；"
                    "source 表示模型来源（Base/Ensemble/Transformer/AutoML）。"
                )
            else:
                self.add_text(f"未找到指标 {indicator} 的跨模型 SHAP 分析表。")
        else:
            self.add_text("未指定 SHAP 关注指标，跳过跨模型分析表。")

"""Table 1 基线特征汇总表生成器。

按 outcome 分组，输出连续变量（均值 ± 标准差）和分类变量（n / %）的汇总表，
并计算组间比较的 P 值，保存为 Excel 文件。
"""

from __future__ import annotations

from pathlib import Path
from typing import Dict, List, Optional, Sequence

import numpy as np
import pandas as pd
from scipy import stats


# 用于 Table 1 的变量角色（排除 id/time/weight/strata/psu/drop/unknown）
TABLE1_ROLES = {"continuous", "binary", "ordinal", "nominal"}


def _format_pvalue(p: float) -> str:
    """格式化 P 值：小于 0.001 显示 <0.001，否则保留 3 位小数。"""
    if np.isnan(p):
        return "NA"
    if p < 0.001:
        return "<0.001"
    return f"{p:.3f}"


def _format_number(n: float, decimals: int = 2) -> str:
    """格式化数字：千分位 + 固定小数位。"""
    return f"{n:,.{decimals}f}"


def _format_pct(n: int, total: int) -> str:
    """格式化计数和百分比。"""
    if total == 0:
        return f"{n:,} (0.00%)"
    pct = n / total * 100
    return f"{n:,} ({pct:.2f}%)"


def _continuous_stats(s: pd.Series) -> Dict[str, float]:
    """计算连续变量的均值、标准差、中位数、Q1、Q3、样本量。"""
    valid = s.dropna()
    return {
        "n": int(valid.shape[0]),
        "mean": float(valid.mean()) if valid.shape[0] > 0 else np.nan,
        "std": float(valid.std()) if valid.shape[0] > 1 else 0.0,
        "median": float(valid.median()) if valid.shape[0] > 0 else np.nan,
        "q1": float(valid.quantile(0.25)) if valid.shape[0] > 0 else np.nan,
        "q3": float(valid.quantile(0.75)) if valid.shape[0] > 0 else np.nan,
    }


def _continuous_pvalue(x: pd.Series, y: pd.Series, non_parametric: bool = False) -> float:
    """两组连续变量比较：默认 t 检验，non_parametric=True 使用 Mann-Whitney U。"""
    a = x.dropna()
    b = y.dropna()
    if a.shape[0] < 2 or b.shape[0] < 2:
        return np.nan
    if non_parametric:
        try:
            _, p = stats.mannwhitneyu(a, b, alternative="two-sided")
        except Exception:
            return np.nan
    else:
        try:
            _, p = stats.ttest_ind(a, b, equal_var=False, nan_policy="omit")
        except Exception:
            return np.nan
    return float(p)


def _categorical_pvalue(group0: pd.Series, group1: pd.Series) -> float:
    """两组分类变量比较：优先卡方检验，期望频数 <5 时退回到 Fisher 精确检验。"""
    try:
        # 合并两组的非缺失类别
        counts0 = group0.value_counts(dropna=True)
        counts1 = group1.value_counts(dropna=True)
        categories = sorted(set(counts0.index) | set(counts1.index))
        if not categories:
            return np.nan
        table = pd.DataFrame(
            {
                "group0": [counts0.get(c, 0) for c in categories],
                "group1": [counts1.get(c, 0) for c in categories],
            },
            index=categories,
        )
        # 卡方检验
        try:
            chi2, p, dof, expected = stats.chi2_contingency(table)
            # 如果任一期望频数 < 5 且为 2x2 列联表，使用 Fisher 精确检验
            if (expected < 5).any() and len(categories) == 2:
                try:
                    _, p = stats.fisher_exact(table)
                except Exception:
                    pass
            return float(p)
        except Exception:
            return np.nan
    except Exception:
        return np.nan


def _build_continuous_row(
    col: str,
    total_stats: Dict[str, float],
    group_stats: List[Dict[str, float]],
    group_labels: List[str],
    pvalue: float,
    show_median: bool = False,
) -> pd.Series:
    """构建连续变量行。"""
    if show_median:
        fmt_total = (
            f"{_format_number(total_stats['mean'])} ± {_format_number(total_stats['std'])} "
            f"[{_format_number(total_stats['median'])} ({_format_number(total_stats['q1'])}-{_format_number(total_stats['q3'])})]"
        )
        fmt_groups = [
            (
                f"{_format_number(gs['mean'])} ± {_format_number(gs['std'])} "
                f"[{_format_number(gs['median'])} ({_format_number(gs['q1'])}-{_format_number(gs['q3'])})]"
            )
            for gs in group_stats
        ]
    else:
        fmt_total = f"{_format_number(total_stats['mean'])} ± {_format_number(total_stats['std'])}"
        fmt_groups = [f"{_format_number(gs['mean'])} ± {_format_number(gs['std'])}" for gs in group_stats]

    data = {
        "Characteristic": col,
        "Total": fmt_total,
        "P value": _format_pvalue(pvalue),
    }
    for label, value in zip(group_labels, fmt_groups):
        data[label] = value
    return pd.Series(data)


def _build_categorical_rows(
    col: str,
    df_total: pd.DataFrame,
    df_groups: List[pd.DataFrame],
    group_labels: List[str],
    pvalue: float,
) -> List[pd.Series]:
    """构建分类变量行：变量名行 + 各类别行。"""
    # 统一按字符串类别处理，避免数值型 0/1 与字符串键不匹配
    total_str = df_total[col].astype(str).replace("nan", pd.NA)
    group_strs = [g[col].astype(str).replace("nan", pd.NA) for g in df_groups]

    # 合并所有组中的非缺失类别，保持原始顺序
    categories = []
    seen = set()
    for sub in [total_str] + group_strs:
        for cat in sub.dropna().unique():
            if cat not in seen:
                categories.append(cat)
                seen.add(cat)

    rows: List[pd.Series] = []
    # 变量名行（带 P 值）
    rows.append(
        pd.Series(
            {
                "Characteristic": f"{col}, n (%)",
                "Total": "",
                "P value": _format_pvalue(pvalue),
            }
            | {label: "" for label in group_labels}
        )
    )

    totals = total_str.value_counts(dropna=True)
    group_counts = [g.value_counts(dropna=True) for g in group_strs]
    for cat in categories:
        row = {
            "Characteristic": f"  {cat}",
            "Total": _format_pct(int(totals.get(cat, 0)), int(df_total.shape[0])),
            "P value": "",
        }
        for label, counts, sub in zip(group_labels, group_counts, df_groups):
            row[label] = _format_pct(int(counts.get(cat, 0)), int(sub.shape[0]))
        rows.append(pd.Series(row))
    return rows


def generate_table1(
    df: pd.DataFrame,
    roles: Dict[str, str],
    outcome_col: Optional[str] = None,
    output_path: Optional[str | Path] = None,
    non_parametric: bool = False,
    show_median: bool = False,
    group_labels: Optional[Sequence[str]] = None,
) -> pd.DataFrame:
    """生成 Table 1 并可选保存为 Excel。

    Parameters
    ----------
    df : pd.DataFrame
        原始数据（预处理前）。
    roles : Dict[str, str]
        变量角色字典。
    outcome_col : str, optional
        结局变量列名；若未指定则自动从 roles 中查找 role == 'outcome' 的列。
    output_path : str | Path, optional
        输出 Excel 路径。若提供则保存。
    non_parametric : bool
        连续变量是否使用 Mann-Whitney U 检验（默认 t 检验）。
    show_median : bool
        连续变量是否同时展示中位数 (IQR)。
    group_labels : Sequence[str], optional
        分组列显示名称；默认按 outcome 实际值生成。

    Returns
    -------
    pd.DataFrame
        格式化后的 Table 1 表格（含 Characteristic, Total, Group... , P value 列）。
    """
    if df is None or df.empty:
        raise ValueError("数据为空，请先在数据页导入数据")

    if outcome_col is None:
        candidates = [c for c, r in roles.items() if r == "outcome"]
        if not candidates:
            raise ValueError("未设置 outcome 角色，请在数据页将结局变量设为 outcome")
        outcome_col = candidates[0]

    if outcome_col not in df.columns:
        raise ValueError(f"结局变量列 {outcome_col} 不在数据中")

    outcome = df[outcome_col].dropna()
    unique_values = sorted(outcome.unique().tolist(), key=lambda x: (str(x), x))
    if len(unique_values) != 2:
        raise ValueError(f"Table 1 目前仅支持二分类结局变量，当前 {outcome_col} 有 {len(unique_values)} 个唯一值：{unique_values}")

    if group_labels is None:
        group_labels = [f"{outcome_col}={v}" for v in unique_values]
    else:
        group_labels = list(group_labels)
        if len(group_labels) != 2:
            raise ValueError("group_labels 必须包含两个标签")

    df_groups = [df[df[outcome_col] == v] for v in unique_values]

    # 收集需要汇总的特征列
    feature_cols = [c for c, r in roles.items() if r in TABLE1_ROLES and c in df.columns and c != outcome_col]
    if not feature_cols:
        raise ValueError("没有可用于 Table 1 的特征列，请将至少一个变量设为 continuous/binary/ordinal/nominal")

    rows: List[pd.Series] = []

    # 样本量行
    n_total = df.shape[0]
    n_groups = [g.shape[0] for g in df_groups]
    rows.append(
        pd.Series(
            {
                "Characteristic": "",
                "Total": f"N = {n_total:,}",
                "P value": "",
            }
            | {label: f"N = {n:,}" for label, n in zip(group_labels, n_groups)}
        )
    )

    for col in feature_cols:
        role = roles.get(col, "continuous")
        if role == "continuous":
            total_stats = _continuous_stats(df[col])
            group_stats = [_continuous_stats(g[col]) for g in df_groups]
            pvalue = _continuous_pvalue(df_groups[0][col], df_groups[1][col], non_parametric=non_parametric)
            rows.append(_build_continuous_row(col, total_stats, group_stats, group_labels, pvalue, show_median))
        else:
            # binary / ordinal / nominal
            pvalue = _categorical_pvalue(df_groups[0][col], df_groups[1][col])
            rows.extend(_build_categorical_rows(col, df, df_groups, group_labels, pvalue))

    # 组装 DataFrame，按指定列序
    columns = ["Characteristic", "Total"] + group_labels + ["P value"]
    result = pd.DataFrame(rows, columns=columns)

    if output_path is not None:
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        _write_excel(result, output_path, group_labels=group_labels)

    return result


def _write_excel(df: pd.DataFrame, output_path: Path, group_labels: List[str]) -> None:
    """将 Table 1 写入 Excel 并进行简单排版。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    from openpyxl.utils.dataframe import dataframe_to_rows

    wb = Workbook()
    ws = wb.active
    ws.title = "Table 1"

    # 写入数据
    for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), 1):
        ws.append(row)

    # 样式
    header_font = Font(bold=True, size=11)
    bold_font = Font(bold=True, size=10)
    normal_font = Font(size=10)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

    # 表头样式
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    # 内容样式
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        is_category_header = row[0].value and str(row[0].value).endswith(", n (%)")
        is_n_row = row[0].value == ""
        for cell in row:
            cell.border = thin_border
            cell.font = bold_font if is_category_header or is_n_row else normal_font
            cell.alignment = Alignment(horizontal="left" if cell.column == 1 else "center", vertical="center", wrap_text=True)

    # 列宽
    ws.column_dimensions["A"].width = 35
    for col_letter in ["B", "C", "D", "E"]:
        ws.column_dimensions[col_letter].width = 22

    # 冻结首行
    ws.freeze_panes = "A2"

    wb.save(output_path)
